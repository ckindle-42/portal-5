"""auto-documents' production suite: sound under the controls gate, gold reproducible, every
deciding fact and reference call declared, and the attainability check able to fail."""

from __future__ import annotations

import base64
import json
import re
from pathlib import Path

import pytest

from tests.wfe import controls
from tests.wfe.suites import build_documents_production as gen

SUITE = Path(__file__).resolve().parents[1] / "suites" / "documents_production.jsonl"
TASKS = {t["id"]: t for t in map(json.loads, SUITE.read_text().splitlines())}


def _kind(prefix: str) -> list[dict]:
    return [t for i, t in TASKS.items() if i.startswith(prefix)]


def test_the_suite_is_sound_and_balanced():
    assert len(TASKS) == 36
    for prefix in ("doc-read-", "doc-create-", "doc-trip-"):
        assert len(_kind(prefix)) == 12
        assert sum(t["screen"] for t in _kind(prefix)) == 4
    assert controls.check_suite(SUITE) == []
    assert controls.schema_failures(next(iter(TASKS.values()))) == []


def test_every_task_is_graded_on_its_answer_or_file_and_on_its_trace():
    for t in TASKS.values():
        kinds = [c["type"] for c in t["checkers"]]
        assert "tool_grounded" in kinds
        assert ("answer_contains" in kinds) != ("office_contains" in kinds), t["id"]
    for t in _kind("doc-trip-"):  # grounded on both the read and the create
        assert sum(c["type"] == "tool_grounded" for c in t["checkers"]) == 2


def test_a_third_of_reads_and_round_trips_are_imperative():
    for prefix in ("doc-read-", "doc-trip-"):
        starts = [t["instruction"].split()[0] for t in _kind(prefix)]
        assert sum(w in ("Read", "Open") for w in starts) >= 4, (prefix, starts)


def test_deciding_facts_and_reference_calls_are_declared_for_every_task():
    gen.build()
    for t in _kind("doc-read-") + _kind("doc-trip-"):
        assert gen.DECIDING_FACTS[t["id"]], t["id"]
    for t in _kind("doc-create-") + _kind("doc-trip-"):
        assert t["id"] in gen.REFERENCE_CALLS, t["id"]


def test_no_enumerated_answer_value_is_left_undefined():
    assert all(not gen.undefined_values(t["instruction"]) for t in TASKS.values())


def test_the_suite_regenerates_identically():
    pytest.importorskip("docx"), pytest.importorskip("openpyxl"), pytest.importorskip("pptx")
    built = "".join(json.dumps(t) + "\n" for t in gen.build())
    assert built == SUITE.read_text()


def test_seeded_fixtures_hold_the_gold(tmp_path):
    pytest.importorskip("docx"), pytest.importorskip("openpyxl"), pytest.importorskip("pptx")
    from tests.wfe.checkers import _office_parts

    for case in gen.read_cases():
        if case["kind"] == "pdf":
            assert str(case["gold"]).encode() in case["raw"]
            continue
        path = tmp_path / case["file"]
        path.write_bytes(case["raw"])
        text, _ = _office_parts(path, case["kind"])
        if case["file"] == "review.pptx":  # speaker notes live outside the shapes
            import pptx

            text = pptx.Presentation(str(path)).slides[1].notes_slide.notes_text_frame.text
        if case["file"] == "budget.xlsx":  # a derived total is not stored in the file
            assert str(case["gold"]) not in text.split()
        else:
            assert str(case["gold"]) in text, case["file"]


def test_attainability_fails_when_a_fact_is_not_exposed_or_the_call_makes_no_file():
    from tests.wfe.e2e import documents_attainability as att

    gen.build()
    read = TASKS["doc-read-00"]

    def silent(tool, args, scope):
        return json.dumps({"success": True, "content": []})

    assert any("fact not exposed" in f for f in att.check_task(read, silent))
    create = TASKS["doc-create-00"]
    assert any("does not satisfy" in f for f in att.check_task(create, silent))


def test_seed_b64_round_trips():
    for t in TASKS.values():
        for rel, b in (t.get("seed_b64") or {}).items():
            assert base64.b64decode(b)[:2] in (b"PK", b"%P"), rel


def test_a_round_trip_instruction_never_states_the_value_it_asks_for():
    gen.build()
    for case, task in zip(gen.trip_cases(), _kind("doc-trip-"), strict=True):
        assert f"{case['label']}: {case['value']}" not in task["instruction"], task["id"]
        bare = task["instruction"].replace(case["file"], "")
        assert not re.search(rf"(?<!\w){re.escape(str(case['value']))}(?!\w)", bare), task["id"]
