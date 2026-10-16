"""Build tests/wfe/suites/documents_production.jsonl, auto-documents' production suite.

Three task kinds, twelve each, over fictional documents built here and seeded into each row's
run directory as binary fixtures (`seed_b64`):

  read   a .docx/.xlsx/.pptx/.pdf holds a fact; the task asks for it (as a question or an
         imperative, one in three). Graded on the answer and on a read_* call whose output
         shows the fact.
  create the task specifies a document's exact content; graded by PARSING the produced file
         (office_contains) and on the create_* call.
  trip   read a seeded xlsx/docx, derive a value (sum, count, lookup, max, product), write it
         into a new document; graded on the parsed output file and on both calls.

Gold is computed here from the same data the fixtures are built from. The fixtures are
byte-reproducible (fixed timestamps, fixed-level deflate). Every fact the gold depends on is in
DECIDING_FACTS (checked live by tests/wfe/e2e/documents_attainability.py), and every create or
trip task has the tool call that produces it in REFERENCE_CALLS (the same check runs it).

    uv run python tests/wfe/suites/build_documents_production.py
"""

from __future__ import annotations

import base64
import datetime as dt
import io
import json
import re
import zipfile
from pathlib import Path
from typing import NamedTuple

OUT = Path(__file__).with_name("documents_production.jsonl")
FIXED = dt.datetime(2030, 1, 1)
ANSWER_FMT = 'Answer only as JSON {"answer": ...}.'
CREATE_OK = r'"success"\s*:\s*true'


class Fact(NamedTuple):
    """One fact the gold depends on, and the read call that must expose it (attainability)."""

    tool: str
    args: dict
    pattern: str
    why: str


DECIDING_FACTS: dict[str, list[Fact]] = {}
#: task id -> (create tool, arguments) that produces the document the task asks for.
REFERENCE_CALLS: dict[str, tuple[str, dict]] = {}
#: task id -> {relative path: base64} fixtures (also in the task's seed_b64).
SEEDS: dict[str, dict[str, str]] = {}

# --------------------------------------------------------------------------
# fixtures: deterministic office and pdf files
# --------------------------------------------------------------------------


def _stable(raw: bytes) -> bytes:
    """The same zip with fixed member dates and level-9 deflate: byte-reproducible on one zlib."""
    out = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(raw)) as src,
        zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst,
    ):
        for info in src.infolist():
            zi = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o600 << 16
            data = src.read(info.filename)
            if info.filename == "docProps/core.xml":  # openpyxl stamps modified=now on save
                data = re.sub(
                    rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>2030-01-01T00:00:00Z", data
                )
            dst.writestr(zi, data, compresslevel=9)
    return out.getvalue()


def _props(cp) -> None:
    cp.author = "Portal"
    cp.last_modified_by = "Portal"
    cp.created = cp.modified = FIXED
    cp.revision = 1


def make_docx(blocks: list[tuple]) -> bytes:
    """blocks: ("title"|"h1"|"h2", text), ("p", text), ("bullets", [text]), ("table", rows)."""
    import docx

    d = docx.Document()
    _props(d.core_properties)
    level = {"title": 0, "h1": 1, "h2": 2}
    for kind, val in blocks:
        if kind in level:
            d.add_heading(val, level=level[kind])
        elif kind == "p":
            d.add_paragraph(val)
        elif kind == "bullets":
            for item in val:
                d.add_paragraph(item, style="List Bullet")
        elif kind == "table":
            t = d.add_table(rows=len(val), cols=len(val[0]))
            t.style = "Table Grid"
            for i, row in enumerate(val):
                for j, cell in enumerate(row):
                    t.cell(i, j).text = str(cell)
    buf = io.BytesIO()
    d.save(buf)
    return _stable(buf.getvalue())


def make_xlsx(sheets: dict[str, list[list]]) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    wb.properties.creator = wb.properties.lastModifiedBy = "Portal"
    wb.properties.created = wb.properties.modified = FIXED
    buf = io.BytesIO()
    wb.save(buf)
    return _stable(buf.getvalue())


def make_pptx(slides: list[tuple], deck_title: str | None = None) -> bytes:
    """slides: (title, content, notes); a deck_title adds the title slide the tool makes."""
    import pptx

    p = pptx.Presentation()
    _props(p.core_properties)
    if deck_title:
        p.slides.add_slide(p.slide_layouts[0]).shapes.title.text = deck_title
    for title, content, *notes in slides:
        s = p.slides.add_slide(p.slide_layouts[1])
        s.shapes.title.text = title
        s.placeholders[1].text = content
        if notes and notes[0]:
            s.notes_slide.notes_text_frame.text = notes[0]
    buf = io.BytesIO()
    p.save(buf)
    return _stable(buf.getvalue())


def make_pdf(pages: list[list[str]]) -> bytes:
    """A minimal text PDF, one page per list of lines. Lines hold no parentheses."""
    n = len(pages)
    objs = ["<</Type/Catalog/Pages 2 0 R>>"]
    objs.append(f"<</Type/Pages/Kids[{' '.join(f'{3 + 2 * i} 0 R' for i in range(n))}]/Count {n}>>")
    font = 3 + 2 * n
    for i, lines in enumerate(pages):
        body = "BT /F1 12 Tf 72 720 Td 16 TL " + " ".join(f"({ln}) Tj T*" for ln in lines) + " ET"
        objs.append(
            f"<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]/Contents {4 + 2 * i} 0 R"
            f"/Resources<</Font<</F1 {font} 0 R>>>>>>"
        )
        objs.append(f"<</Length {len(body)}>>\nstream\n{body}\nendstream")
    objs.append("<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>")
    out, offs = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offs.append(len(out))
        out += f"{i} 0 obj\n{o}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offs)
    return (
        out + f"trailer\n<</Size {len(objs) + 1}/Root 1 0 R>>\nstartxref\n{xref}\n%%EOF\n".encode()
    )


def b64(raw: bytes) -> str:
    return base64.b64encode(raw).decode()


# --------------------------------------------------------------------------
# answer and evidence patterns
# --------------------------------------------------------------------------


def answer_pattern(gold: object) -> str:
    if isinstance(gold, str):
        return rf'"answer"\s*:\s*"{re.escape(gold)}"'
    return rf'"answer"\s*:\s*"?{gold}"?(?![\d.])'


def evidence(value: object) -> str:
    """A pattern for a value in a read tool's JSON output (cells are stringified)."""
    return rf"(?<![\d.]){value}(?!\d)" if isinstance(value, int) else re.escape(str(value))


def undefined_values(instruction: str) -> list[str]:
    """Enumerated answer values (`... is one of A, B, C: ...`) with no definition (none here)."""
    from tests.wfe.suites.build_compliance_production import undefined_values as uv

    return uv(instruction)


# --------------------------------------------------------------------------
# read tasks
# --------------------------------------------------------------------------

READ_TOOL = {
    "docx": "read_word_document",
    "xlsx": "read_excel",
    "pptx": "read_powerpoint",
    "pdf": "read_pdf",
}


def read_cases() -> list[dict]:
    """Twelve reads. `direct` finishes "In <file>, ...?"; `ask` finishes "Read <file> and tell
    me ..."; `gold` is computed from the data the fixture is built from."""
    vendor = [["Clause", "Days"], ["Notice", 37], ["Renewal", 90], ["Cure", 14]]
    staff = [
        ["Name", "Team", "Desk"],
        ["Ines Calder", "Billing", 114],
        ["Tobias Wren", "Ops", 207],
        ["Mara Quill", "Billing", 118],
        ["Dov Harnett", "Ops", 203],
    ]
    stock = {
        "East": [["Part", "Qty"], ["bracket", 12], ["anvil", 31]],
        "West": [["Part", "Qty"], ["bracket", 47], ["anvil", 9]],
    }
    budget = [["Line", "Cost"], ["Rent", 1200], ["Power", 340], ["Water", 85], ["Tools", 610]]
    deck = [
        ("Kickoff", "Agenda and goals"),
        ("Rollout", "Owner: Ilse Mott"),
        ("Budget", "Approved spend: 48000"),
    ]
    notes = [
        ("Overview", "Team intro", "Open with the safety moment."),
        ("Costs", "Summary table", "The contingency budget is 7200 euros."),
        ("Close", "Next steps", ""),
    ]
    steps = [
        ("Intake", "Receive the request"),
        ("Triage", "Rank by urgency"),
        ("Review", "Check the draft"),
        ("Archive", "File the record"),
    ]
    cell = lambda rows, key, col: next(r[col] for r in rows if r[0] == key)  # noqa: E731
    return [
        {
            "kind": "docx",
            "file": "memo.docx",
            "raw": make_docx(
                [
                    ("h1", "Harbor Lease Memo"),
                    ("p", "The escrow agent is Pell Varga."),
                    ("p", "The notary is Odile Rask."),
                ]
            ),
            "direct": "who is named as the escrow agent",
            "ask": "who is named as the escrow agent",
            "gold": "Pell Varga",
            "facts": ["Pell Varga"],
        },
        {
            "kind": "docx",
            "file": "terms.docx",
            "raw": make_docx([("h1", "Vendor Terms"), ("table", vendor)]),
            "direct": "how many days is the Renewal period",
            "ask": "how many days the Renewal period is",
            "gold": cell(vendor[1:], "Renewal", 1),
            "facts": ["Renewal", 90],
        },
        {
            "kind": "docx",
            "file": "survey.docx",
            "raw": make_docx(
                [
                    ("h1", "Site Survey"),
                    ("h2", "North Yard"),
                    ("p", "The gate code is 4471."),
                    ("h2", "South Yard"),
                    ("p", "The gate code is 5826."),
                ]
            ),
            "direct": "what is the gate code for the South Yard",
            "ask": "the gate code for the South Yard",
            "gold": 5826,
            "facts": ["5826", "South Yard"],
        },
        {
            "kind": "xlsx",
            "file": "staff.xlsx",
            "raw": make_xlsx({"Staff": staff}),
            "direct": "which desk is Mara Quill assigned to",
            "ask": "which desk Mara Quill is assigned to",
            "gold": cell(staff[1:], "Mara Quill", 2),
            "facts": ["Mara Quill", 118],
        },
        {
            "kind": "xlsx",
            "file": "stock.xlsx",
            "raw": make_xlsx(stock),
            "direct": "how many brackets are listed on the West sheet",
            "ask": "how many brackets are listed on the West sheet",
            "gold": cell(stock["West"][1:], "bracket", 1),
            "facts": ["West", 47],
        },
        {
            "kind": "xlsx",
            "file": "budget.xlsx",
            "raw": make_xlsx({"Budget": budget}),
            "direct": "what is the total of the Cost column",
            "ask": "the total of the Cost column",
            "gold": sum(r[1] for r in budget[1:]),
            "facts": ["Rent", 1200, 340, 85, 610],
        },
        {
            "kind": "pptx",
            "file": "launch.pptx",
            "raw": make_pptx([(t, c) for t, c in deck]),
            "direct": "who is the owner named on the Rollout slide",
            "ask": "who the owner named on the Rollout slide is",
            "gold": "Ilse Mott",
            "facts": ["Rollout", "Ilse Mott"],
        },
        {
            "kind": "pptx",
            "file": "review.pptx",
            "raw": make_pptx(notes),
            "direct": "what contingency budget do the speaker notes of the Costs slide state? Give the number only",
            "ask": "what contingency budget the speaker notes of the Costs slide state, as a number only",
            "gold": 7200,
            "facts": ["Costs", 7200],
        },
        {
            "kind": "pptx",
            "file": "process.pptx",
            "raw": make_pptx(steps),
            "direct": "what is the title of the third slide",
            "ask": "the title of the third slide",
            "gold": steps[2][0],
            "facts": ["Intake", "Review", "Archive"],
        },
        {
            "kind": "pdf",
            "file": "incident.pdf",
            "raw": make_pdf([["Incident Record", "Case ref OR-7731.", "The vault code is 5521."]]),
            "direct": "what is the vault code",
            "ask": "the vault code",
            "gold": 5521,
            "facts": ["5521"],
        },
        {
            "kind": "pdf",
            "file": "inspection.pdf",
            "raw": make_pdf(
                [
                    ["Annual Inspection", "Page one lists the scope."],
                    ["Findings", "The inspector is Sunniva Beck.", "Next visit 2031-04-18."],
                ]
            ),
            "direct": "on which date is the next visit scheduled",
            "ask": "the date of the next visit, in the form YYYY-MM-DD",
            "gold": "2031-04-18",
            "facts": ["2031-04-18", "Sunniva Beck"],
        },
        {
            "kind": "pdf",
            "file": "suppliers.pdf",
            "raw": make_pdf(
                [
                    ["Supplier List", "Primary: Fenn Metals", "Backup: Alder Cast"],
                    ["Terms", "Payment due in 45 days."],
                ]
            ),
            "direct": "in how many days is payment due",
            "ask": "how many days until payment is due",
            "gold": 45,
            "facts": ["45 days"],
        },
    ]


def read_tasks() -> list[dict]:
    tasks = []
    for i, c in enumerate(read_cases()):
        tool = READ_TOOL[c["kind"]]
        tid = f"doc-read-{i:02d}"
        # one in three is an imperative: the suite measures how users phrase it
        text = (
            f"Read {c['file']} and tell me {c['ask']}."
            if i % 3 == 0
            else f"In {c['file']}, {c['direct']}?"
        )
        gold = c["gold"]
        ev = [evidence(f) for f in c["facts"][:2]]
        DECIDING_FACTS[tid] = [
            Fact(tool, {"file_path": c["file"]}, evidence(f), f"the fixture holds {f}")
            for f in c["facts"]
        ]
        seed = {c["file"]: b64(c["raw"])}
        SEEDS[tid] = seed
        trace = [
            {
                "name": tool,
                "args": {"file_path": c["file"]},
                "output": json.dumps({"success": True, "facts": [str(f) for f in c["facts"]]}),
                "error": None,
            }
        ]
        wrong = "none" if isinstance(gold, str) else gold + 7
        tasks.append(
            {
                "id": tid,
                "lane": "documents",
                "agentic": True,
                "screen": i % 3 == 0,
                "instruction": f"{text} {ANSWER_FMT}",
                "requires_tools": [tool],
                "seed_b64": seed,
                "checkers": [
                    {"type": "answer_contains", "patterns": [answer_pattern(gold)]},
                    {"type": "tool_grounded", "tools": [tool], "evidence": ev},
                ],
                "controls": {
                    "pass": [{"answer": json.dumps({"answer": gold}), "tool_calls": trace}],
                    "fail": [
                        json.dumps({"answer": gold}),
                        {"answer": json.dumps({"answer": wrong}), "tool_calls": trace},
                    ],
                },
            }
        )
    return tasks


# --------------------------------------------------------------------------
# create tasks
# --------------------------------------------------------------------------


def md(blocks: list[tuple]) -> str:
    """The markdown create_word_document renders: headings, bullets, paragraphs, tables."""
    out = []
    for kind, val in blocks:
        if kind == "h1":
            out.append(f"# {val}")
        elif kind == "h2":
            out.append(f"## {val}")
        elif kind == "p":
            out.append(val)
        elif kind == "bullets":
            out.append("\n".join(f"- {v}" for v in val))
        elif kind == "table":
            head = "| " + " | ".join(map(str, val[0])) + " |"
            sep = "|" + "---|" * len(val[0])
            body = ["| " + " | ".join(map(str, r)) + " |" for r in val[1:]]
            out.append("\n".join([head, sep, *body]))
    return "\n\n".join(out)


def describe_docx(blocks: list[tuple]) -> str:
    parts = []
    for kind, val in blocks:
        if kind == "h1":
            parts.append(f'a top-level heading "{val}"')
        elif kind == "h2":
            parts.append(f'a second-level heading "{val}"')
        elif kind == "p":
            parts.append(f'a paragraph reading "{val}"')
        elif kind == "bullets":
            parts.append("a bulleted list with the items " + ", ".join(f'"{v}"' for v in val))
        elif kind == "table":
            parts.append(
                "a table whose header row is "
                + " | ".join(map(str, val[0]))
                + " and whose rows are "
                + "; ".join(" | ".join(map(str, r)) for r in val[1:])
            )
    return ", then ".join(parts)


def docx_strings(blocks: list[tuple]) -> tuple[list[str], list[str]]:
    texts, heads = [], []
    for kind, val in blocks:
        if kind in ("h1", "h2"):
            heads.append(val)
            texts.append(val)
        elif kind == "p":
            texts.append(val)
        elif kind == "bullets":
            texts += val
        elif kind == "table":
            texts += [str(c) for r in val for c in r]
    return texts, heads


CREATE_DOCX = [
    (
        "Orchard Safety Checklist",
        [
            ("h1", "Before Picking"),
            ("bullets", ["Check ladders", "Wear gloves", "Log the weather"]),
        ],
    ),
    (
        "Quarterly Greenhouse Plan",
        [
            ("h1", "Goals"),
            ("p", "Raise tomato yield by twelve percent."),
            ("h2", "Risks"),
            ("p", "Late frost may delay planting."),
        ],
    ),
    (
        "Pump Maintenance Roster",
        [
            ("h1", "Roster"),
            (
                "table",
                [
                    ["Pump", "Technician", "Interval"],
                    ["P-14", "Corin Alder", "30 days"],
                    ["P-22", "Yusra Fenn", "45 days"],
                ],
            ),
        ],
    ),
    (
        "Visitor Log Template",
        [
            ("table", [["Name", "Arrival"], ["Lucan Thorne", "09:15"], ["Petra Voss", "10:40"]]),
            ("p", "Keep this log at the front desk."),
        ],
    ),
]
CREATE_XLSX = [
    (
        "Seed Order Tracker",
        {
            "Orders": [
                ["Seed", "Qty", "Supplier"],
                ["fennel", 40, "Brack Farm"],
                ["sorrel", 25, "Holm Seeds"],
                ["chervil", 60, "Brack Farm"],
            ]
        },
    ),
    (
        "Regional Rainfall",
        {
            "North": [["Month", "Millimetres"], ["March", 82], ["April", 64]],
            "South": [["Month", "Millimetres"], ["March", 51], ["April", 47]],
        },
    ),
    (
        "Lab Calibration",
        {
            "Sheet1": [
                ["Probe", "Offset", "Checked"],
                ["A1", 3, "yes"],
                ["B2", 7, "no"],
                ["C3", 2, "yes"],
                ["D4", 9, "no"],
            ]
        },
    ),
    (
        "Crew Rota",
        {
            "Week 12": [
                ["Crew", "Shift"],
                ["Harbour", "Early"],
                ["Quay", "Late"],
                ["Depot", "Night"],
            ]
        },
    ),
]
CREATE_PPTX = [
    (
        "Dock Safety Briefing",
        [
            ("Hazards", "Wet floors"),
            ("Controls", "Barriers and signage"),
            ("Contacts", "Duty officer Mara Quill"),
        ],
    ),
    (
        "Harvest Review",
        [
            ("Yields", "Tomatoes up twelve percent"),
            ("Costs", "Packaging rose by four percent"),
            ("Plans", "Add a second greenhouse"),
            ("Thanks", "Pickers and drivers"),
        ],
    ),
    ("Onboarding Week", [("Monday", "Badge and laptop"), ("Wednesday", "Security training")]),
    (
        "Fleet Update",
        [
            ("Vans", "Six in service"),
            ("Trucks", "Two awaiting parts"),
            ("Outlook", "Replace one truck in spring"),
        ],
    ),
]


def _create_trace(tool: str, name: str) -> list[dict]:
    out = json.dumps({"success": True, "filename": name, "message": f"created: {name}"})
    return [{"name": tool, "args": {"title": "x"}, "output": out, "error": None}]


def flat_file(kind: str, strings: list[str]) -> bytes:
    """A near-miss: a file of this kind holding some of the strings (no structure)."""
    if kind == "docx":
        return make_docx([("p", s) for s in strings])
    if kind == "xlsx":
        return make_xlsx({"Sheet1": [[s] for s in strings]})
    return make_pptx([("Notes", s) for s in strings])


def _file_task(
    tid: str,
    screen: bool,
    instruction: str,
    kind: str,
    create_tool: str,
    spec: dict,
    gold_file: bytes,
    reads: list[dict] | None = None,
    seed: dict | None = None,
    near_miss: bytes | None = None,
) -> dict:
    """One produced-file task. `spec` is the office_contains spec (expect + structure)."""
    trace = _create_trace(create_tool, f"out.{kind}")
    for r in reads or []:
        trace = [
            {
                "name": r["tool"],
                "args": {"file_path": r["file"]},
                "output": json.dumps({"success": True, "facts": r["facts"]}),
                "error": None,
            },
            *trace,
        ]
    checkers = [
        {"type": "office_contains", "kind": kind, **spec},
        *(
            {"type": "tool_grounded", "tools": [r["tool"]], "evidence": r["evidence"]}
            for r in reads or []
        ),
        {"type": "tool_grounded", "tools": [create_tool], "evidence": [CREATE_OK]},
    ]
    if not near_miss:  # drop a string the rest of the file cannot supply
        drop = next(
            e
            for e in reversed(spec["expect"])
            if e not in " ".join(x for x in spec["expect"] if x != e)
        )
        near_miss = flat_file(kind, [x for x in spec["expect"] if x != drop])
    near = near_miss
    t = {
        "id": tid,
        "lane": "documents",
        "agentic": True,
        "screen": screen,
        "instruction": instruction,
        "requires_tools": [*(r["tool"] for r in reads or []), create_tool],
        "checkers": checkers,
        "controls": {
            "pass": [
                {
                    "answer": "Done.",
                    "tool_calls": trace,
                    "files_b64": {f"out/gold.{kind}": b64(gold_file)},
                }
            ],
            "fail": [
                {
                    "answer": "Done.",
                    "tool_calls": [],
                    "files_b64": {f"out/gold.{kind}": b64(gold_file)},
                },
                {
                    "answer": "Done.",
                    "tool_calls": trace,
                    "files_b64": {f"out/near.{kind}": b64(near)},
                },
                {"answer": "Done.", "tool_calls": trace},
            ],
        },
    }
    if seed:
        t["seed_b64"] = seed
    return t


def create_tasks() -> list[dict]:
    tasks = []
    for i, (title, blocks) in enumerate(CREATE_DOCX):
        tid = f"doc-create-{i:02d}"
        texts, heads = docx_strings(blocks)
        REFERENCE_CALLS[tid] = ("create_word_document", {"title": title, "content": md(blocks)})
        tasks.append(
            _file_task(
                tid,
                True,
                f'Create a Word document titled "{title}" containing {describe_docx(blocks)}.',
                "docx",
                "create_word_document",
                {"expect": [title, *texts], "headings": [title, *heads]},
                make_docx([("title", title), *blocks]),
            )
        )
    for i, (title, sheets) in enumerate(CREATE_XLSX):
        tid = f"doc-create-{4 + i:02d}"
        strings = [str(c) for rows in sheets.values() for r in rows for c in r]
        if len(sheets) == 1:
            ((name, rows),) = sheets.items()
            args = {"title": title, "data": rows, "sheet_name": name}
            where = f'a sheet named "{name}"'
            body = (
                f"whose first row is the header {' | '.join(map(str, rows[0]))} and whose rows are "
                + "; ".join(" | ".join(map(str, r)) for r in rows[1:])
            )
        else:
            args = {
                "title": title,
                "sheets": [{"name": n, "headers": r[0], "rows": r[1:]} for n, r in sheets.items()],
            }
            where = "two sheets, " + " and ".join(f'"{n}"' for n in sheets)
            body = (
                "each with the header row "
                + " | ".join(map(str, next(iter(sheets.values()))[0]))
                + "; "
                + "; ".join(
                    f'"{n}" has the rows ' + "; ".join(" | ".join(map(str, r)) for r in rows[1:])
                    for n, rows in sheets.items()
                )
            )
        REFERENCE_CALLS[tid] = ("create_excel", args)
        tasks.append(
            _file_task(
                tid,
                True,
                f'Create an Excel workbook titled "{title}" with {where}, {body}.',
                "xlsx",
                "create_excel",
                {"expect": strings, "sheets": list(sheets)},
                make_xlsx(sheets),
            )
        )
    for i, (title, slides) in enumerate(CREATE_PPTX):
        tid = f"doc-create-{8 + i:02d}"
        REFERENCE_CALLS[tid] = (
            "create_powerpoint",
            {"title": title, "slides": [{"title": t, "content": c} for t, c in slides]},
        )
        listing = "; ".join(
            f'slide {n} is titled "{t}" with the text "{c}"' for n, (t, c) in enumerate(slides, 1)
        )
        tasks.append(
            _file_task(
                tid,
                True,
                f'Create a PowerPoint deck titled "{title}" with {len(slides)} content slides: {listing}.',
                "pptx",
                "create_powerpoint",
                {"expect": [title, *(c for _, c in slides)], "titles": [t for t, _ in slides]},
                make_pptx(slides, deck_title=title),
            )
        )
    for t in tasks:
        t["screen"] = int(t["id"][-2:]) % 4 == 0
    return tasks


# --------------------------------------------------------------------------
# round trip: read, derive, write
# --------------------------------------------------------------------------

PARTS = ["gasket", "valve", "flange", "hose", "clamp", "bearing", "coupler"]
JOBS = ["survey", "wiring", "plaster", "paint", "glaze", "tile", "grout"]


def parts_rows(i: int) -> list[list]:
    """Six distinct-unit rows: Item, Units, Price (no ties, so max and lookup are unique)."""
    return [
        [PARTS[(i + j) % 7], 5 + (i * 13 + j * 17) % 40 + j, 2 + (i * 5 + j * 3) % 25]
        for j in range(6)
    ]


def jobs_rows(i: int) -> list[list]:
    return [[JOBS[(i + j) % 7], 8 + (i * 11 + j * 19) % 50 + j] for j in range(6)]


def trip_cases() -> list[dict]:
    """Twelve cases: source kind, the data, the operation, the label and value, the target."""
    cases = []

    def case(src, file, rows, ask, label, value, target, title, sheet=None, extra=None):
        cases.append(
            {
                "src": src,
                "file": file,
                "rows": rows,
                "ask": ask,
                "label": label,
                "value": value,
                "target": target,
                "title": title,
                "sheet": sheet,
                "extra": extra,
            }
        )

    r = [parts_rows(i) for i in range(6)]
    case(
        "xlsx",
        "orders.xlsx",
        r[0],
        "add up the Units column",
        "Total units",
        sum(x[1] for x in r[0]),
        "docx",
        "Units Report",
    )
    case(
        "xlsx",
        "bulk.xlsx",
        r[1],
        "count the rows with more than 20 Units",
        "Large orders",
        sum(x[1] > 20 for x in r[1]),
        "xlsx",
        "Bulk Orders",
    )
    case(
        "xlsx",
        "prices.xlsx",
        r[2],
        f"find the Price of {r[2][3][0]}",
        f"Price of {r[2][3][0]}",
        r[2][3][2],
        "pptx",
        "Price Check",
    )
    case(
        "xlsx",
        "volume.xlsx",
        r[3],
        "find the Item with the most Units",
        "Top item",
        max(r[3], key=lambda x: x[1])[0],
        "docx",
        "Top Seller",
    )
    ops = [["Line", "Cost"], *[[JOBS[j], 90 + j * 37 % 61 + j] for j in range(5)]]
    case(
        "xlsx",
        "costs.xlsx",
        ops,
        "add up the Cost column on the Ops sheet",
        "Ops total",
        sum(x[1] for x in ops[1:]),
        "xlsx",
        "Ops Cost",
        sheet="Ops",
        extra={"Archive": [["Line", "Cost"], ["old", 999]]},
    )
    status = [
        ["Ticket", "Status"],
        ["T-11", "open"],
        ["T-12", "closed"],
        ["T-13", "open"],
        ["T-14", "open"],
        ["T-15", "closed"],
    ]
    case(
        "xlsx",
        "tickets.xlsx",
        status,
        "count the tickets whose Status is open",
        "Open tickets",
        sum(x[1] == "open" for x in status[1:]),
        "pptx",
        "Ticket Status",
    )
    j = [jobs_rows(i) for i in range(5)]
    case(
        "docx",
        "plan.docx",
        j[0],
        "add up the Hours column of the table",
        "Total hours",
        sum(x[1] for x in j[0]),
        "docx",
        "Hours Report",
    )
    case(
        "docx",
        "roster.docx",
        j[1],
        f"find the Hours for {j[1][2][0]}",
        f"Hours for {j[1][2][0]}",
        j[1][2][1],
        "xlsx",
        "Hours Lookup",
    )
    case(
        "docx",
        "backlog.docx",
        j[2],
        "count the tasks with more than 30 Hours",
        "Long tasks",
        sum(x[1] > 30 for x in j[2]),
        "pptx",
        "Long Tasks",
    )
    case(
        "docx",
        "capacity.docx",
        None,
        "multiply the number of sites by the capacity per site",
        "Total capacity",
        6 * 120,
        "docx",
        "Capacity Report",
    )
    case(
        "docx",
        "effort.docx",
        j[3],
        "find the Task with the most Hours",
        "Heaviest task",
        max(j[3], key=lambda x: x[1])[0],
        "xlsx",
        "Heaviest Task",
    )
    case(
        "docx",
        "sprint.docx",
        j[4],
        "add up the Hours column of the table",
        "Sprint hours",
        sum(x[1] for x in j[4]),
        "pptx",
        "Sprint Hours",
    )
    return cases


def trip_target(target: str, title: str, label: str, value: object) -> dict:
    """What a round trip writes: the instruction clause, the reference call, the gold file, the
    office_contains spec and a near-miss file holding the right label but no value."""
    line = f"{label}: {value}"
    if target == "docx":
        return {
            "tgt": f'create a Word document titled "{title}" whose body is the single line "{line}" with the value you worked out',
            "tool": "create_word_document",
            "args": {"title": title, "content": line},
            "gold": make_docx([("title", title), ("p", line)]),
            "spec": {"expect": [line], "headings": [title]},
            "wrong": make_docx([("title", title), ("p", f"{label}: nothing")]),
        }
    if target == "xlsx":
        rows = [["Metric", "Value"], [label, value]]
        return {
            "tgt": f'create an Excel workbook titled "{title}" with the header row Metric | Value and one data row {label} | <the value you worked out>',
            "tool": "create_excel",
            "args": {"title": title, "data": rows},
            "gold": make_xlsx({"Sheet1": rows}),
            "spec": {"expect": ["Metric", "Value", label, str(value)]},
            "wrong": make_xlsx({"Sheet1": [["Metric", "Value"], [label, "n/a"]]}),
        }
    return {
        "tgt": f'create a PowerPoint deck titled "{title}" with one content slide titled "Summary" whose text is "{line}" with the value you worked out',
        "tool": "create_powerpoint",
        "args": {"title": title, "slides": [{"title": "Summary", "content": line}]},
        "gold": make_pptx([("Summary", line)], deck_title=title),
        "spec": {"expect": [line], "titles": ["Summary"]},
        "wrong": make_pptx([("Summary", f"{label}: nothing")], deck_title=title),
    }


def trip_tasks() -> list[dict]:
    tasks = []
    for i, c in enumerate(trip_cases()):
        tid = f"doc-trip-{i:02d}"
        src, tool = c["src"], READ_TOOL[c["src"]]
        if src == "xlsx":
            if c["file"] == "costs.xlsx":
                sheets = {c["sheet"]: c["rows"], **c["extra"]}
                body_rows = c["rows"][1:]
            elif c["file"] == "tickets.xlsx":
                sheets = {"Tickets": c["rows"]}
                body_rows = c["rows"][1:]
            else:
                sheets = {"Orders": [["Item", "Units", "Price"], *c["rows"]]}
                body_rows = c["rows"]
            raw = make_xlsx(sheets)
        elif c["rows"] is None:
            body_rows = []
            raw = make_docx(
                [
                    ("h1", "Capacity Notes"),
                    ("p", "The plant runs 6 sites."),
                    ("p", "Each site holds 120 units."),
                ]
            )
        else:
            body_rows = c["rows"]
            raw = make_docx([("h1", "Work Plan"), ("table", [["Task", "Hours"], *c["rows"]])])
        seed = {c["file"]: b64(raw)}
        SEEDS[tid] = seed
        facts = (
            [str(x) for row in body_rows for x in row] if body_rows else ["6 sites", "120 units"]
        )
        DECIDING_FACTS[tid] = [
            Fact(
                tool,
                {"file_path": c["file"]},
                evidence(f if not f.isdigit() else int(f)),
                "an input of the derived value",
            )
            for f in facts
        ]
        ev = [evidence(facts[0]), evidence(facts[-1])] if body_rows else [evidence("6 sites")]
        if body_rows and facts[0].isdigit():
            ev = [evidence(int(facts[0])), evidence(int(facts[-1]))]
        label, value, title = c["label"], c["value"], c["title"]
        sheet_note = f" on the {c['sheet']} sheet" if c["sheet"] else ""
        verb = f"Read {c['file']}, {c['ask']}" if i % 3 == 0 else f"Open {c['file']} and {c['ask']}"
        t = trip_target(c["target"], title, label, value)
        REFERENCE_CALLS[tid] = (t["tool"], t["args"])
        tgt, gold, spec, wrong, ctool = t["tgt"], t["gold"], t["spec"], t["wrong"], t["tool"]
        instruction = f"{verb}{sheet_note}, then {tgt}."
        tasks.append(
            _file_task(
                tid,
                False,  # the screen tier is set in build()
                instruction,
                c["target"],
                ctool,
                spec,
                gold,
                reads=[
                    {
                        "tool": tool,
                        "file": c["file"],
                        "evidence": ev,
                        "facts": facts,
                    }
                ],
                seed=seed,
                near_miss=wrong,
            )
        )
    return tasks


def build() -> list[dict]:
    DECIDING_FACTS.clear()
    REFERENCE_CALLS.clear()
    SEEDS.clear()
    tasks = [*read_tasks(), *create_tasks(), *trip_tasks()]
    for t in tasks:
        if bad := undefined_values(t["instruction"]):
            raise SystemExit(f"{t['id']}: undefined answer values {bad}")
    # four per kind in the screen tier
    for kind in ("read", "create", "trip"):
        of_kind = [t for t in tasks if t["id"].startswith(f"doc-{kind}-")]
        for n, t in enumerate(of_kind):
            t["screen"] = n % 3 == 0
    tasks.sort(key=lambda t: not t["screen"])
    return tasks


def main() -> None:
    tasks = build()
    OUT.write_text("".join(json.dumps(t) + "\n" for t in tasks))
    print(f"wrote {OUT.name}: {len(tasks)} tasks, {sum(t['screen'] for t in tasks)} screen")


if __name__ == "__main__":
    main()
