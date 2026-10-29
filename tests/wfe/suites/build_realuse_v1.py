"""Build deterministic fixtures referenced by the runnable realuse_v1 suite."""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import subprocess
import sys
import tempfile
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

import pymupdf
from docx import Document
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[3]
SUITE = ROOT / "tests/wfe/suites/realuse_v1.jsonl"
DRAFT = ROOT / "tests/wfe/suites/realuse_v1.draft.jsonl"
MANIFEST = Path("tests/fixtures/realuse_v1/manifest.sha256")
PDF_METADATA = {
    "title": "",
    "author": "",
    "subject": "",
    "keywords": "",
    "creator": "realuse fixture builder",
    "producer": "PyMuPDF",
    "creationDate": "D:20000101000000Z",
    "modDate": "D:20000101000000Z",
}
TEXT_EXPECTATIONS = {
    "ru-v1-001": ("Timeline slipped by 3 days", "2400", "890", "340", "180", "210", "4020"),
    "ru-v1-002": (
        "The Kitchen",
        "Sweet Basil",
        "Bin 707",
        "45 minutes",
        "2 hours",
        "3 hours",
        "90 minutes",
        "tire chains",
    ),
    "ru-v1-003": ("30 min max", "blocker", "Lisbon", "Tokyo"),
    "ru-v1-004": (
        "Tuesday",
        "Wednesday",
        "Thursday",
        "10am-2pm ET",
        "$300",
        "$500",
        "does not discuss international remote work",
    ),
    "ru-v1-006": ("September", "October", "Priya", "frozen", "August"),
    "ru-v1-009": ("ERROR missing config.json", "READY"),
    "ru-v1-010": ("urgent", "priority"),
    "ru-v1-012": ("load_data", "json.load", "value", "42"),
    "ru-v1-013": (
        "From: alice@example.com",
        "Subject: Meeting Tomorrow",
        "From: bob@example.com",
        "See you at noon.",
    ),
    "ru-v1-014": ("8 percent", "1.05"),
    "ru-v1-018": (
        "1 Scope",
        "2 Payment",
        "3 Termination",
        "4 Liability",
        "5 Governing law",
        "Net 30",
        "$500",
    ),
    "ru-v1-019": ("Acme Supply", "$1,240.00", "[illegible]"),
    "ru-v1-028": ("chickadee", "once a week", "twice a week", "bread"),
    "ru-v1-029": ("pie", "not a cake", "Saturday 9:00am"),
    "ru-v1-034": ("cfg['db']", "cfg['mode']", "'db': 'shop'", "'mode': 'live'"),
    "ru-v1-036": ("total // people",),
    "ru-v1-037": ("prices['Tea']",),
    "ru-v1-038": ("pages * 2",),
    "ru-v1-039": ("'shelf'",),
    "ru-v1-040": ("sorted(rows)",),
    "ru-v1-041": ("sum(readings)",),
    "ru-v1-042": ("parts[2]", "decaf"),
    "ru-v1-043": ("'N '",),
    "ru-v1-044": ("hours", "22"),
    "ru-v1-045": ("range(1, len(days))",),
    "ru-v1-046": ("TODO",),
    "ru-v1-047": ("strptime", "2026/03/05"),
    "ru-v1-048": ("renames the elapsed column to seconds", "row['elapsed']"),
    "ru-v1-056": ("31", "18"),
    "ru-v1-064": ("Dana", "crane", "Leo", "October 28"),
    "ru-v1-065": ("knee", "jumps"),
    "ru-v1-070": ("October 20", "December 15", "gym"),
}
SEED_EXPECTATIONS = {
    "ru-v1-009": ("Thursday 2pm", "Friday"),
    "ru-v1-021": ("Monday through Wednesday", "No equipment stipend"),
    "ru-v1-023": ("55 minutes", "72 minutes", "2022"),
    "ru-v1-024": ("500 ml", "3 minutes", "4 minutes", "40 dollars", "Fuel is included"),
    "ru-v1-030": ("AWS", "GCP", "$15000 per month"),
}
IMAGE_EXPECTATIONS = {
    "ru-v1-025": ("Portal 5 — Open Source Intelligence Layer", "0, 51, 102", "24x24"),
    "ru-v1-026": ("P1", "P2", "P3", "P4", "P5", "Q3 Team"),
    "ru-v1-033": ("line", "bar", "pie", "LOW STOCK", "200, 30, 30"),
}


#: (fixture name, broken snippet, fixed snippet, "crash"|"value", broken marker, fixed marker).
REPAIR_SPECS = {
    "ru-v1-036": ("tip_split.py", "total // people", "total / people", "value", "23", "23.5"),
    "ru-v1-037": ("orders.py", "'Tea'", "'tea'", "crash", "KeyError", "3"),
    "ru-v1-038": (
        "reading_time.py",
        'print(minutes, "hours")',
        'print(minutes // 60, "hours")',
        "value",
        "180 hours",
        "3 hours",
    ),
    "ru-v1-039": ("backstock.py", " if item['place'] == 'shelf'", "", "value", "12", "19"),
    "ru-v1-040": (
        "leaderboard.py",
        "sorted(rows)",
        "sorted(rows, key=lambda row: -row['score'])",
        "crash",
        "TypeError",
        "Ann",
    ),
    "ru-v1-041": (
        "readings.py",
        " / len(readings)",
        " / len(readings) if readings else 0",
        "crash",
        "ZeroDivisionError",
        "average=0",
    ),
    "ru-v1-042": (
        "columns.py",
        "print(parts[2])",
        "print(parts[2] if len(parts) > 2 else '-')",
        "crash",
        "IndexError",
        "decaf",
    ),
    "ru-v1-043": ("tally.py", "vote == 'N'", "vote.strip() == 'N'", "value", "noes=0", "noes=2"),
    "ru-v1-044": (
        "payroll.py",
        "member['hours'] * 22",
        "(member['hours'] or 0) * 22",
        "crash",
        "TypeError",
        "Mara 0",
    ),
    "ru-v1-045": ("week.py", "range(1, len(days))", "range(len(days))", "value", "1 Tue", "0 Mon"),
    "ru-v1-047": (
        "dates.py",
        "'%m/%d/%Y'",
        "'%Y/%m/%d'",
        "crash",
        "ValueError",
        "2026-03-05 00:00:00",
    ),
    "ru-v1-048": ("metrics.py", "row['elapsed']", "row['seconds']", "crash", "KeyError", "90"),
}


def _jsonl(path: Path) -> list[dict[str, Any]]:
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line]


def _spec_path(task_id: str, spec: dict[str, Any]) -> Path | None:
    place = spec.get("place", "")
    upload_prefix = "/workspace/uploads/realuse_v1/"
    if place.startswith(upload_prefix):
        return Path("tests/fixtures/realuse_v1") / place.removeprefix(upload_prefix)
    if place.startswith("fixtures/realuse_v1/"):
        return Path(place.split(" under ", maxsplit=1)[0])
    if place.startswith("eval fixture") or place.startswith("workspace upload"):
        return Path("tests/fixtures/realuse_v1") / task_id / spec["name"]
    return None


def _referenced_specs(
    suite_rows: list[dict[str, Any]], draft_rows: list[dict[str, Any]]
) -> dict[Path, dict[str, Any]]:
    drafts = {row["id"]: row for row in draft_rows}
    found: dict[Path, dict[str, Any]] = {}
    for row in suite_rows:
        task_id = row["id"]
        specs = [
            item
            for item in drafts[task_id].get("fixtures", [])
            if item.get("format") not in {"none", "memory", "kb"}
        ]
        for fixture in row.get("fixtures", []):
            rel = Path(fixture)
            if rel.is_absolute() or ".." in rel.parts:
                raise ValueError(f"unsafe suite fixture path: {fixture}")
            candidates = [item for item in specs if item.get("name") == rel.name]
            if len(candidates) != 1:
                raise ValueError(
                    f"{task_id}: expected one v2 spec for {fixture}, found {len(candidates)}"
                )
            if _spec_path(task_id, candidates[0]) != rel:
                raise ValueError(f"{task_id}: v2 place does not match suite path {fixture}")
            if rel in found:
                raise ValueError(f"duplicate suite fixture path: {fixture}")
            found[rel] = candidates[0]
    return found


def _stable_docx(spec: dict[str, Any]) -> bytes:
    document = Document()
    for paragraph in spec["paragraphs"]:
        document.add_paragraph(paragraph)
    properties = document.core_properties
    properties.author = ""
    properties.last_modified_by = ""
    properties.created = properties.modified = datetime(2000, 1, 1)
    raw = io.BytesIO()
    document.save(raw)
    stable = io.BytesIO()
    with (
        ZipFile(io.BytesIO(raw.getvalue())) as source,
        ZipFile(stable, "w", ZIP_DEFLATED) as target,
    ):
        for name in source.namelist():
            info = ZipInfo(name, (2000, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            info.create_system = 3
            target.writestr(info, source.read(name))
    return stable.getvalue()


def _pdf(text: str) -> bytes:
    document = pymupdf.open()
    page = document.new_page(width=612, height=792)
    spare = page.insert_textbox(
        pymupdf.Rect(48, 48, 564, 744), text, fontname="helv", fontsize=12, lineheight=1.35
    )
    if spare < 0:
        raise ValueError("fixture PDF text does not fit on one page")
    document.set_metadata(PDF_METADATA)
    result = io.BytesIO()
    document.save(result, garbage=4, deflate=True, no_new_id=True)
    document.close()
    return result.getvalue()


def _font(size: int) -> ImageFont.FreeTypeFont:
    candidates = ("Arial.ttf", "DejaVuSans.ttf", "/System/Library/Fonts/Supplemental/Arial.ttf")
    for candidate in candidates:
        try:
            return ImageFont.truetype(candidate, size)
        except OSError:
            continue
    raise OSError("no supported TrueType font is installed")


def _jpeg(image: Image.Image) -> bytes:
    output = io.BytesIO()
    image.save(output, format="JPEG", quality=95, subsampling=0, optimize=False, progressive=False)
    return output.getvalue()


def _image(task_id: str) -> bytes:
    if task_id == "ru-v1-025":
        image = Image.new("RGB", (640, 360), (0, 51, 102))
        draw = ImageDraw.Draw(image)
        draw.rectangle((16, 16, 39, 39), fill=(255, 220, 0))
        title = "Portal 5 — Open Source Intelligence Layer"
        font = _font(22)
        box = draw.textbbox((0, 0), title, font=font)
        draw.text(((640 - (box[2] - box[0])) / 2, 170), title, fill="white", font=font)
    elif task_id == "ru-v1-026":
        image = Image.new("RGB", (800, 500), "white")
        draw = ImageDraw.Draw(image)
        for index in range(5):
            left = 45 + index * 145
            draw.rectangle((left, 100, left + 115, 360), outline=(35, 55, 80), width=3)
            draw.text((left + 38, 112), f"P{index + 1}", fill="black", font=_font(22))
            if index < 4:
                draw.rectangle((left + 12, 155, left + 50, 178), fill=(255, 220, 0))
        draw.text((320, 410), "Q3 Team", fill="black", font=_font(28))
    elif task_id == "ru-v1-033":
        image = Image.new("RGB", (800, 400), (180, 190, 200))
        draw = ImageDraw.Draw(image)
        draw.line((65, 210, 105, 170, 145, 195, 185, 135), fill=(30, 80, 150), width=6)
        draw.text((110, 235), "line", fill="black", font=_font(20))
        for index, height in enumerate((70, 115, 90)):
            left = 300 + index * 42
            draw.rectangle((left, 210 - height, left + 28, 210), fill=(50, 100, 160))
        draw.text((305, 235), "bar", fill="black", font=_font(20))
        draw.pieslice((500, 120, 640, 260), 25, 300, fill=(230, 160, 50), outline="black")
        draw.text((540, 275), "pie", fill="black", font=_font(20))
        draw.rectangle((535, 35, 775, 95), fill=(200, 30, 30))
        draw.text((552, 50), "LOW STOCK", fill="white", font=_font(25))
    else:
        raise ValueError(f"unsupported image task: {task_id}")
    return _jpeg(image)


def _fixture_bytes(task_id: str, spec: dict[str, Any]) -> bytes:
    fmt = spec["format"]
    if fmt == "docx":
        return _stable_docx(spec)
    if fmt == "pdf":
        return _pdf(spec["text"])
    if fmt == "jpg":
        return _image(task_id)
    if fmt in {"csv", "json", "py", "txt"}:
        return spec["text"].encode("utf-8")
    raise ValueError(f"unsupported fixture format: {fmt}")


def _rendered_text(rel: Path, raw: bytes) -> str:
    if rel.suffix == ".pdf":
        document = pymupdf.open(stream=raw, filetype="pdf")
        text = "\n".join(page.get_text() for page in document)
        document.close()
        return text
    if rel.suffix == ".docx":
        return "\n".join(paragraph.text for paragraph in Document(io.BytesIO(raw)).paragraphs)
    return raw.decode("utf-8")


def _run(code: str) -> subprocess.CompletedProcess[str]:
    return subprocess.run([sys.executable, "-c", code], capture_output=True, text=True, check=False)


def _csv_rows(raw: bytes) -> list[dict[str, str]]:
    return list(csv.DictReader(io.StringIO(raw.decode("utf-8"))))


def _check_bug_tasks(files: dict[Path, bytes]) -> None:
    with tempfile.TemporaryDirectory() as temp:
        local = Path(temp) / "customers.csv"
        local.write_bytes(files[Path("tests/fixtures/realuse_v1/ru-v1-005/customers.csv")])
        source = files[Path("tests/fixtures/realuse_v1/ru-v1-005/count_active.py")].decode()
        source = source.replace("/workspace/uploads/realuse_v1/ru-v1-005/customers.csv", str(local))
        before = _run(source)
        after = _run(source.replace("rows = rows[:-1]", "rows = rows"))
        assert before.returncode == 0 and "active=4" in before.stdout
        assert after.returncode == 0 and "active=5" in after.stdout

    calc = files[Path("tests/fixtures/realuse_v1/ru-v1-011/calc.py")].decode()
    before = _run(calc + "\nprint(total(5))\n")
    after = _run(calc.replace("range(1, n)", "range(1, n + 1)") + "\nprint(total(5))\n")
    assert before.returncode == after.returncode == 0
    assert before.stdout.strip() == "10" and after.stdout.strip() == "15"

    discount = files[Path("tests/fixtures/realuse_v1/ru-v1-031/discount.py")].decode()
    before = _run(discount + "\nprint(final_price(100))\n")
    after = _run("TEN = 10\n" + discount + "\nprint(final_price(100))\n")
    assert before.returncode != 0 and "NameError" in before.stderr
    assert after.returncode == 0 and after.stdout.strip() == "90"

    quote = files[Path("tests/fixtures/realuse_v1/ru-v1-014/quote.py")].decode()
    before = _run(quote + "\nprint(total(100))\n")
    after = _run(quote.replace("1.05", "1.08") + "\nprint(total(100))\n")
    assert before.returncode == after.returncode == 0
    assert before.stdout.strip() == "105.0" and after.stdout.strip() == "108.0"


def _check_repair_tasks(files: dict[Path, bytes]) -> None:
    for task_id, (name, old, new, shape, broken, fixed_out) in REPAIR_SPECS.items():
        source = files[Path(f"tests/fixtures/realuse_v1/{task_id}/{name}")].decode()
        before = _run(source)
        after = _run(source.replace(old, new))
        if shape == "crash":
            assert before.returncode != 0 and broken in before.stderr
        else:
            assert before.returncode == 0 and broken in before.stdout
        assert after.returncode == 0 and fixed_out in after.stdout


def _check_expand_counts(files: dict[Path, bytes]) -> None:
    notes = files[Path("tests/fixtures/realuse_v1/ru-v1-046/notes.txt")].decode().splitlines()
    assert len(notes) == 8 and sum("TODO" in line for line in notes) == 3
    temps = [
        int(line)
        for line in files[Path("tests/fixtures/realuse_v1/ru-v1-056/temperatures.txt")]
        .decode()
        .split()
    ]
    assert max(temps) == 31 and min(temps) == 18


def _check_csv_tasks(files: dict[Path, bytes]) -> None:
    customers = _csv_rows(files[Path("tests/fixtures/realuse_v1/ru-v1-005/customers.csv")])
    assert len(customers) == 8
    assert sum(row["status"] == "active" for row in customers) == 5

    orders = _csv_rows(files[Path("tests/fixtures/realuse_v1/ru-v1-008/orders.csv")])
    regions = Counter()
    weeks = Counter()
    for row in orders:
        regions[row["region"]] += int(row["orders"])
        weeks[row["week"]] += int(row["orders"])
    assert regions == {"north": 76, "south": 57}
    assert min(weeks, key=weeks.get) == "4" and weeks["4"] == 21

    ratings = _csv_rows(files[Path("tests/fixtures/realuse_v1/ru-v1-022/survey.csv")])
    north = [int(row["rating"]) for row in ratings if row["region"] == "north"]
    low = Counter(row["product_line"] for row in ratings if int(row["rating"]) <= 2)
    assert len(ratings) == 14 and sum(north) / len(north) == 2.0
    assert low == {"stove": 2, "lantern": 1}

    adopted = _csv_rows(files[Path("tests/fixtures/realuse_v1/ru-v1-035/adopted.csv")])
    adopted_count = sum(row["adopted"] == "yes" for row in adopted)
    assert len(adopted) == 25 and adopted_count == 8
    assert adopted_count * 100 / len(adopted) == 32 and 40 - 32 == 8


def _check_document_tasks(files: dict[Path, bytes]) -> None:
    drive = _rendered_text(
        Path("drive_notes.docx"),
        files[Path("tests/fixtures/realuse_v1/ru-v1-002/drive_notes.docx")],
    )
    assert (
        "Denver to Boulder: 45 minutes" in drive
        and "Sweet Basil" in drive
        and "tire chains" in drive
    )
    notes = [
        files[Path(f"tests/fixtures/realuse_v1/ru-v1-006/note_0{index}.docx")]
        for index in range(1, 5)
    ]
    assert len(notes) == 4
    note_text = " ".join(_rendered_text(Path("note.docx"), raw) for raw in notes)
    assert all(word in note_text for word in ("September", "October", "Priya", "August"))


def _check_script_tasks(files: dict[Path, bytes]) -> None:
    report = files[Path("tests/fixtures/realuse_v1/ru-v1-010/report.txt")].decode().splitlines()
    assert "urgent" in report[1] and "priority" in report[4] and "urgent" in report[6]
    assert sum("urgent" in line for line in report) == 2
    assert sum("priority" in line for line in report) == 1

    email = files[Path("tests/fixtures/realuse_v1/ru-v1-013/email_nosubject.txt")].decode()
    assert "From: bob@example.com" in email and "Subject:" not in email

    prep = files[Path("tests/fixtures/realuse_v1/ru-v1-009/prep.py")].decode()
    before = _run(prep)
    live = _run(f"exec({prep!r}, {{'mode': 'live'}})")
    assert before.returncode == live.returncode == 0
    assert before.stdout.strip() == "ERROR missing config.json" and live.stdout.strip() == "READY"


def _check_pdf_tasks(files: dict[Path, bytes]) -> None:
    for rel, raw in files.items():
        if rel.suffix == ".pdf":
            assert len(pymupdf.open(stream=raw, filetype="pdf")) == 1


def _check_code_tasks(files: dict[Path, bytes]) -> None:
    source = files[Path("fixtures/realuse_v1/ru-v1-012/app.py")].decode()
    data = json.loads(files[Path("fixtures/realuse_v1/ru-v1-012/data.json")])
    assert (
        data["value"] == 42
        and "def load_data(path):" in source
        and "json.load(handle)['value']" in source
    )
    namespace: dict[str, Any] = {"json": json}
    with tempfile.TemporaryDirectory() as temp:
        present = Path(temp) / "data.json"
        missing = Path(temp) / "missing.json"
        present.write_text(json.dumps(data), encoding="utf-8")
        exec(compile(source, "app.py", "exec"), namespace)
        assert namespace["load_data"](present) == 42
        try:
            namespace["load_data"](missing)
        except FileNotFoundError as error:
            assert not str(error).startswith("missing:")
        else:
            raise AssertionError("unfixed load_data unexpectedly accepted a missing file")
        fixed = source.replace(
            "    with open(path, encoding='utf-8') as handle:\n        return json.load(handle)['value']",
            "    try:\n        with open(path, encoding='utf-8') as handle:\n            return json.load(handle)['value']\n"
            "    except FileNotFoundError as error:\n        raise FileNotFoundError(f'missing: {path}') from error",
        )
        fixed_namespace: dict[str, Any] = {"json": json}
        exec(compile(fixed, "app.py", "exec"), fixed_namespace)
        assert fixed_namespace["load_data"](present) == 42
        try:
            fixed_namespace["load_data"](missing)
        except FileNotFoundError as error:
            assert "missing:" in str(error)
        else:
            raise AssertionError("reference fix did not reject missing data")

    entry = files[Path("fixtures/realuse_v1/ru-v1-034/entry_point.py")].decode()
    cfg = files[Path("fixtures/realuse_v1/ru-v1-034/config_reader.py")].decode()
    helper = files[Path("fixtures/realuse_v1/ru-v1-034/helper_b.py")].decode()
    fixed_entry = entry.replace("return cfg['db']", "return cfg.get('db', 'none')")
    assert "return cfg['db']" in entry and "raw.get('path'" in cfg and "cfg['mode']" in helper
    for value, expected in (({"db": "shop"}, "shop"), ({}, "none")):
        module = type("ConfigReader", (), {"load_config": staticmethod(lambda value=value: value)})
        namespace = {"config_reader": module}
        exec(
            compile(fixed_entry.replace("import config_reader\n", ""), "entry_point.py", "exec"),
            namespace,
        )
        assert namespace["main"]() == expected


def _check_line_and_image_tasks(files: dict[Path, bytes]) -> None:
    line_files = {
        "a.py": 3,
        "b.py": 2,
        "c.py": 4,
        "tests/t1.py": 5,
        "tests/t2.py": 1,
    }
    base = Path("tests/fixtures/realuse_v1/ru-v1-015")
    counts = {name: files[base / name].decode().count("\n") for name in line_files}
    assert counts == line_files and sum(counts.values()) == 15
    assert sum(count for name, count in counts.items() if not name.startswith("tests/")) == 9

    app = files[Path("tests/fixtures/realuse_v1/ru-v1-016/app.py")].decode().splitlines()
    assert app[17] == "CACHE = []" and app[41] == '    handle = open(path, encoding="utf-8")'
    assert "handle.close()" in "\n".join(app) and "CACHE.append(text)" in "\n".join(app)

    for task_id, size in (
        ("ru-v1-025", (640, 360)),
        ("ru-v1-026", (800, 500)),
        ("ru-v1-033", (800, 400)),
    ):
        rel = next(path for path in files if path.parent.name == task_id)
        image = Image.open(io.BytesIO(files[rel])).convert("RGB")
        assert image.size == size
        if task_id == "ru-v1-025":
            yellow = image.getpixel((27, 27))
            assert all(abs(a - b) <= 12 for a, b in zip(yellow, (255, 220, 0), strict=True))
            background = image.getpixel((600, 300))
            assert all(abs(a - b) <= 12 for a, b in zip(background, (0, 51, 102), strict=True))
        elif task_id == "ru-v1-026":
            badges = [image.getpixel((left + 25, 165)) for left in (45, 190, 335, 480, 625)]
            assert sum(pixel[0] > 220 and pixel[1] > 180 and pixel[2] < 80 for pixel in badges) == 4
        else:
            banner = image.getpixel((540, 45))
            assert banner[0] > 150 and banner[1] < 70 and banner[2] < 70


def _verify_task_values(files: dict[Path, bytes]) -> None:
    grouped: dict[str, list[tuple[Path, bytes]]] = defaultdict(list)
    for rel, raw in files.items():
        if len(rel.parts) > 2 and rel.parts[0] == "tests" and rel.parts[1] == "fixtures":
            task_id = rel.parts[3]
        elif len(rel.parts) > 2 and rel.parts[0] == "fixtures":
            task_id = rel.parts[2]
        else:
            continue
        grouped[task_id].append((rel, raw))
    for task_id, values in TEXT_EXPECTATIONS.items():
        text = " ".join(_rendered_text(rel, raw) for rel, raw in grouped[task_id])
        missing = [value for value in values if value.casefold() not in text.casefold()]
        if missing:
            raise AssertionError(f"{task_id}: checker values missing from fixtures: {missing}")
    _check_csv_tasks(files)
    _check_document_tasks(files)
    _check_pdf_tasks(files)
    _check_code_tasks(files)
    _check_script_tasks(files)
    _check_line_and_image_tasks(files)
    _check_bug_tasks(files)
    _check_repair_tasks(files)
    _check_expand_counts(files)


def _check_nonfile_seeds(
    suite_rows: list[dict[str, Any]], draft_rows: list[dict[str, Any]]
) -> None:
    drafts = {row["id"]: row for row in draft_rows}
    for row in suite_rows:
        specs = drafts[row["id"]].get("fixtures", [])
        memories = [item["seed"] for item in specs if item.get("format") == "memory"]
        knowledge = [item["text"] for item in specs if item.get("format") == "kb"]
        if memories != row.get("memory_seed", []):
            raise AssertionError(f"{row['id']}: runnable memory seed differs from the v2 draft")
        seed_values = list((row.get("kb_seed") or {}).values())
        if knowledge != seed_values:
            raise AssertionError(f"{row['id']}: runnable KB seed differs from the v2 draft")
        combined = " ".join(memories + knowledge)
        missing = [
            value
            for value in SEED_EXPECTATIONS.get(row["id"], ())
            if value.casefold() not in combined.casefold()
        ]
        if missing:
            raise AssertionError(
                f"{row['id']}: checker values missing from runnable non-file seeds: {missing}"
            )


def _check_image_specs(draft_rows: list[dict[str, Any]]) -> None:
    rows = {row["id"]: row for row in draft_rows}
    for task_id, expected in IMAGE_EXPECTATIONS.items():
        images = [item for item in rows[task_id].get("fixtures", []) if item.get("format") == "jpg"]
        if len(images) != 1:
            raise AssertionError(f"{task_id}: expected one image spec")
        spec = images[0].get("generate", "")
        missing = [value for value in expected if value.casefold() not in spec.casefold()]
        if missing:
            raise AssertionError(f"{task_id}: image checker values missing from v2 spec: {missing}")


def _fixture_bytes_by_path(suite_path: Path, draft_path: Path) -> dict[Path, bytes]:
    specs = _referenced_specs(_jsonl(suite_path), _jsonl(draft_path))
    return {rel: _fixture_bytes(rel.parts[-2], spec) for rel, spec in specs.items()}


def build(
    destination_root: Path = ROOT,
    suite_path: Path = SUITE,
    draft_path: Path = DRAFT,
) -> dict[Path, str]:
    files = _fixture_bytes_by_path(suite_path, draft_path)
    hashes: dict[Path, str] = {}
    for rel, raw in sorted(files.items()):
        target = destination_root / rel
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(raw)
        hashes[rel] = hashlib.sha256(raw).hexdigest()
    manifest_text = "".join(
        f"{digest}  {rel.as_posix()}\n" for rel, digest in sorted(hashes.items())
    )
    manifest_path = destination_root / MANIFEST
    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    manifest_path.write_text(manifest_text, encoding="utf-8")
    return hashes


def self_check(
    destination_root: Path = ROOT,
    suite_path: Path = SUITE,
    draft_path: Path = DRAFT,
) -> dict[Path, str]:
    suite_rows = _jsonl(suite_path)
    draft_rows = _jsonl(draft_path)
    expected = _referenced_specs(suite_rows, draft_rows)
    expected_bytes = {rel: _fixture_bytes(rel.parts[-2], spec) for rel, spec in expected.items()}
    actual_paths = set()
    for folder in (Path("tests/fixtures/realuse_v1"), Path("fixtures/realuse_v1")):
        actual_paths.update(
            path.relative_to(destination_root)
            for path in (destination_root / folder).rglob("*")
            if path.is_file() and path != destination_root / MANIFEST
        )
    if actual_paths != set(expected_bytes):
        raise AssertionError(
            f"fixture path mismatch: missing={sorted(set(expected_bytes) - actual_paths)} extra={sorted(actual_paths - set(expected_bytes))}"
        )
    actual = {rel: (destination_root / rel).read_bytes() for rel in expected_bytes}
    for rel, raw in expected_bytes.items():
        if actual[rel] != raw:
            raise AssertionError(f"fixture differs from its v2 spec: {rel}")
    hashes = {rel: hashlib.sha256(raw).hexdigest() for rel, raw in actual.items()}
    manifest_text = "".join(
        f"{digest}  {rel.as_posix()}\n" for rel, digest in sorted(hashes.items())
    )
    if (destination_root / MANIFEST).read_text(encoding="utf-8") != manifest_text:
        raise AssertionError("fixture manifest does not match generated files")
    _verify_task_values(actual)
    _check_nonfile_seeds(suite_rows, draft_rows)
    _check_image_specs(draft_rows)
    return hashes


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--check", action="store_true", help="verify generated fixtures and manifest"
    )
    parser.add_argument(
        "--self-check",
        action="store_true",
        help="build fixtures, then verify checker values and bug fixes",
    )
    args = parser.parse_args()
    hashes = self_check() if args.check else build()
    if args.self_check:
        hashes = self_check()
    print(f"realuse_v1 fixtures: {len(hashes)} paths, manifest {MANIFEST}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
